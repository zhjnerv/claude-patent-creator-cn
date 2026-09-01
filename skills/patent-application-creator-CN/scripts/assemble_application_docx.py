#!/usr/bin/env python3
"""按 Word 模板组装中国发明专利申请文件，并生成原生 Word 公式。

输入目录约定包含：权利要求书.md、说明书.md、说明书摘要.md、说明书附图.md
以及说明书附图目录。输出沿用模板的样式、分节、页眉、页脚、行号和页码设置。

公式支持两种来源：
1. 推荐：Markdown 行内 ``$...$``、独立 ``$$...$$`` 或 ``math`` 代码块；
2. 兼容：独立成段且含等号的线性公式，例如 ``ED_CT=Q_CT×C_CT(P,a)``。

含公式时必须在 Windows 上通过 Microsoft Word 的 OMaths.Add/BuildUp 转换为
可编辑公式；禁止把普通字符排版冒充为公式。默认交付只执行 DOCX 结构化校验；
仅当用户明确要求视觉检查时，才由 Word 导出 PDF 并调用 pdftoppm 生成逐页 PNG。
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
import sys
import tempfile
from collections.abc import Iterable
from contextlib import suppress
from copy import deepcopy
from dataclasses import dataclass
from datetime import datetime
from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt
except ImportError as exc:  # pragma: no cover - 环境错误路径
    raise SystemExit("缺少 python-docx；请在执行环境中安装 python-docx 后重试。") from exc

try:
    from lxml import etree
except ImportError as exc:  # pragma: no cover - 项目依赖缺失路径
    raise SystemExit("缺少 lxml；请先安装项目依赖。") from exc

try:
    from PIL import Image
except ImportError as exc:  # pragma: no cover - 项目依赖缺失路径
    raise SystemExit("缺少 Pillow；请先安装项目依赖。") from exc


EXPECTED_HEADERS = ["权利要求书", "说明书", "说明书附图", "说明书摘要", "摘要附图"]
REQUIRED_PARAGRAPH_STYLES = ["Normal (Web)", "Title", "Heading 1", "正文2", "附图图号"]
REQUIRED_CHARACTER_STYLE = "Strong"  # 中文 Word 界面显示为“要点”
REPORT_SCHEMA = "cn-patent-docx-assembly/v2"

CJK_RE = re.compile(r"[\u3400-\u9fff]")
DISPLAY_FORMULA_RE = re.compile(r"=")
INLINE_EXPLICIT_RE = re.compile(r"\$([^$\n]+)\$")
INLINE_EQUATION_RE = re.compile(
    r"(?<![A-Za-z0-9_])"
    r"[A-Za-z][A-Za-z0-9_]*(?:_(?:[A-Za-z0-9]+|\([A-Za-z0-9+\-]+\)))?"
    r"(?:\([^，。；\s]*\))?="
    r"[A-Za-z0-9_()+\-*/×^.,≈]+"
)
IDENTIFIER_RE = re.compile(r"[A-Za-z]+(?:_(?:[A-Za-z0-9]+|\([A-Za-z0-9+\-]+\)))*")


@dataclass(frozen=True)
class SpecItem:
    kind: str
    text: str


@dataclass(frozen=True)
class FigureSpec:
    number: int
    title: str
    path: Path


@dataclass(frozen=True)
class MathSpec:
    marker: str
    linear: str


class MathRegistry:
    def __init__(self) -> None:
        self.specs: list[MathSpec] = []

    def add(self, linear: str) -> str:
        value = normalize_formula(linear)
        if not value:
            raise ValueError("公式内容为空")
        marker = f"[[EQ{len(self.specs) + 1:04d}]]"
        self.specs.append(MathSpec(marker=marker, linear=value))
        return marker


def read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8").replace("\r\n", "\n")


def file_sha256(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest()


def markdown_blocks(path: Path) -> list[str]:
    text = read_text(path).strip()
    return [
        re.sub(r"\s*\n\s*", " ", block.strip())
        for block in re.split(r"\n\s*\n", text)
        if block.strip()
    ]


def parse_claims(path: Path) -> list[str]:
    blocks = markdown_blocks(path)
    if not blocks or blocks[0] != "# 权利要求书":
        raise ValueError("权利要求书缺少预期标题“# 权利要求书”")
    claims: list[str] = []
    numbers: list[int] = []
    for block in blocks[1:]:
        match = re.match(r"^(\d+)\.\s*(.+)$", block, flags=re.S)
        if not match:
            raise ValueError(f"无法识别权利要求编号：{block[:80]}")
        numbers.append(int(match.group(1)))
        claims.append(match.group(2).strip())
    if numbers != list(range(1, len(claims) + 1)):
        raise ValueError(f"权利要求编号不连续：{numbers}")
    return claims


def normalize_formula(value: str) -> str:
    text = value.strip()
    if text.startswith("$$") and text.endswith("$$"):
        text = text[2:-2].strip()
    if text.startswith("```math") and text.endswith("```"):
        text = text[len("```math") : -3].strip()
    return text.removesuffix("。").strip()


def looks_like_display_formula(value: str) -> bool:
    text = normalize_formula(value)
    if not DISPLAY_FORMULA_RE.search(text):
        return False
    if len(text) > 500 or CJK_RE.search(text):
        return False
    return bool(re.search(r"[_^×*/()]", text))


def classify_spec_block(value: str) -> SpecItem:
    text = value.strip()
    if text.startswith("$$") and text.endswith("$$"):
        return SpecItem("formula", normalize_formula(text))
    if text.startswith("```math") and text.endswith("```"):
        return SpecItem("formula", normalize_formula(text))
    if looks_like_display_formula(text):
        return SpecItem("formula", normalize_formula(text))
    return SpecItem("body", text)


def parse_specification(path: Path) -> list[SpecItem]:
    items: list[SpecItem] = []
    pending: list[str] = []
    math_fence: list[str] | None = None

    def flush() -> None:
        if pending:
            items.append(classify_spec_block(" ".join(line.strip() for line in pending)))
            pending.clear()

    for raw in read_text(path).splitlines():
        line = raw.strip()
        if math_fence is not None:
            if line == "```":
                items.append(SpecItem("formula", normalize_formula(" ".join(math_fence))))
                math_fence = None
            else:
                math_fence.append(line)
            continue
        if line == "```math":
            flush()
            math_fence = []
            continue
        if not line:
            flush()
            continue
        if line.startswith("# "):
            flush()
            items.append(SpecItem("title", line[2:].strip()))
        elif line.startswith("## "):
            flush()
            items.append(SpecItem("heading", line[3:].strip()))
        elif line.startswith("### "):
            flush()
            items.append(SpecItem("heading", line[4:].strip()))
        else:
            pending.append(line)
    if math_fence is not None:
        raise ValueError("说明书存在未闭合的 ```math 代码块")
    flush()
    if not items or items[0].kind != "title":
        raise ValueError("说明书缺少一级标题形式的发明名称")
    return items


def parse_abstract(path: Path) -> tuple[str, int]:
    blocks = markdown_blocks(path)
    if not blocks or not blocks[0].startswith("# "):
        raise ValueError("说明书摘要缺少一级标题")
    figure_number: int | None = None
    body: list[str] = []
    for block in blocks[1:]:
        match = re.fullmatch(r"摘要附图：图(\d+)[。.]?", block)
        if match:
            figure_number = int(match.group(1))
        else:
            body.append(block)
    if len(body) != 1:
        raise ValueError(f"预期一个摘要正文段落，实际为 {len(body)} 个")
    if figure_number is None:
        raise ValueError("说明书摘要未指定“摘要附图：图N。”")
    return body[0], figure_number


def resolve_figure_path(index_path: Path, raw_target: str) -> Path:
    target = Path(raw_target)
    candidate = (index_path.parent / target).resolve()
    if candidate.suffix.lower() == ".svg":
        png = candidate.with_suffix(".png")
        if png.is_file():
            candidate = png
    if not candidate.is_file():
        raise FileNotFoundError(f"附图文件不存在：{candidate}")
    if candidate.suffix.lower() not in {".png", ".jpg", ".jpeg"}:
        raise ValueError(f"Word 组装仅接受 PNG/JPEG 发布图：{candidate}")
    return candidate


def parse_figures(path: Path) -> list[FigureSpec]:
    heading_re = re.compile(r"^##\s+图(\d+)\s+(.+)$")
    image_re = re.compile(r"^!\[[^]]*]\(([^)]+)\)$")
    current: tuple[int, str] | None = None
    figures: list[FigureSpec] = []
    for raw in read_text(path).splitlines():
        line = raw.strip()
        heading = heading_re.match(line)
        if heading:
            current = (int(heading.group(1)), heading.group(2).strip())
            continue
        image = image_re.match(line)
        if image and current is not None:
            figures.append(
                FigureSpec(
                    number=current[0],
                    title=current[1],
                    path=resolve_figure_path(path, image.group(1)),
                )
            )
            current = None
    numbers = [figure.number for figure in figures]
    if not figures or numbers != list(range(1, len(figures) + 1)):
        raise ValueError(f"说明书附图编号不连续：{numbers}")
    return figures


def extract_function_calls(formula: str) -> set[str]:
    calls: set[str] = set()
    for match in re.finditer(r"[A-Za-z][A-Za-z0-9_]*\(", formula):
        start = match.start()
        depth = 0
        for index in range(match.end() - 1, len(formula)):
            char = formula[index]
            if char == "(":
                depth += 1
            elif char == ")":
                depth -= 1
                if depth == 0:
                    calls.add(formula[start : index + 1])
                    break
    return calls


def collect_math_symbols(items: Iterable[SpecItem]) -> tuple[str, ...]:
    symbols: set[str] = set()
    for item in items:
        if item.kind != "formula":
            continue
        symbols.update(IDENTIFIER_RE.findall(item.text))
        symbols.update(extract_function_calls(item.text))
    # 常量和函数名也可作为公式对象；数字与运算符不单独提取。
    return tuple(sorted((value for value in symbols if value), key=len, reverse=True))


def tokenize_inline_math(text: str, symbols: tuple[str, ...]) -> list[tuple[bool, str]]:
    """把正文拆成普通文字与需要转为 OMath 的线性表达式。"""
    explicit: list[tuple[int, int, str]] = [
        (match.start(), match.end(), match.group(1)) for match in INLINE_EXPLICIT_RE.finditer(text)
    ]
    symbol_pattern = None
    if symbols:
        symbol_pattern = re.compile(
            r"(?<![A-Za-z0-9_])(?:"
            + "|".join(re.escape(value) for value in symbols)
            + r")(?![A-Za-z0-9_])"
        )

    spans: list[tuple[int, int, str]] = list(explicit)
    protected = [(start, end) for start, end, _ in explicit]

    def overlaps(start: int, end: int) -> bool:
        return any(start < other_end and end > other_start for other_start, other_end in protected)

    for match in INLINE_EQUATION_RE.finditer(text):
        if not overlaps(match.start(), match.end()):
            spans.append((match.start(), match.end(), match.group(0)))
            protected.append((match.start(), match.end()))
    if symbol_pattern is not None:
        for match in symbol_pattern.finditer(text):
            if not overlaps(match.start(), match.end()):
                spans.append((match.start(), match.end(), match.group(0)))
                protected.append((match.start(), match.end()))

    spans.sort(key=lambda item: item[0])
    result: list[tuple[bool, str]] = []
    cursor = 0
    for start, end, formula in spans:
        if start < cursor:
            continue
        if start > cursor:
            result.append((False, text[cursor:start]))
        result.append((True, formula))
        cursor = end
    if cursor < len(text):
        result.append((False, text[cursor:]))
    return result or [(False, text)]


def validate_template(doc: Document) -> list:
    if len(doc.sections) != 5:
        raise ValueError(f"输出模板必须包含5个分节，实际为 {len(doc.sections)}")
    missing = [name for name in REQUIRED_PARAGRAPH_STYLES if name not in doc.styles]
    if missing:
        raise ValueError(f"输出模板缺少段落样式：{missing}")
    if REQUIRED_CHARACTER_STYLE not in doc.styles:
        raise ValueError("输出模板缺少“要点”字符样式（OOXML/英文名 Strong）")
    headers = [
        section.header.paragraphs[0].text.strip() if section.header.paragraphs else ""
        for section in doc.sections
    ]
    if headers != EXPECTED_HEADERS:
        raise ValueError(f"输出模板页眉顺序错误：{headers}")
    first = doc.paragraphs[0]._p.pPr if doc.paragraphs else None
    if first is None or first.numPr is None:
        raise ValueError("输出模板首段必须提供权利要求自动编号格式")
    return [deepcopy(section._sectPr) for section in doc.sections]


def clear_body_keep_final_sectpr(doc: Document, final_sectpr) -> None:
    body = doc._body._element
    for child in list(body):
        body.remove(child)
    body.append(deepcopy(final_sectpr))


def copy_paragraph_properties(paragraph, source_ppr) -> None:
    existing = paragraph._p.pPr
    if existing is not None:
        paragraph._p.remove(existing)
    paragraph._p.insert(0, deepcopy(source_ppr))


def copy_run_properties(run, source_rpr) -> None:
    if source_rpr is None:
        return
    existing = run._r.rPr
    if existing is not None:
        run._r.remove(existing)
    run._r.insert(0, deepcopy(source_rpr))


def set_east_asia_font(run, font_name: str) -> None:
    run.font.name = font_name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = rpr._add_rFonts()
    rfonts.set(qn("w:eastAsia"), font_name)
    rfonts.set(qn("w:ascii"), font_name)
    rfonts.set(qn("w:hAnsi"), font_name)


def end_section(paragraph, sectpr) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    existing = ppr.sectPr
    if existing is not None:
        ppr.remove(existing)
    ppr.append(deepcopy(sectpr))


def add_claim(doc: Document, text: str, claim_ppr, claim_rpr):
    paragraph = doc.add_paragraph()
    copy_paragraph_properties(paragraph, claim_ppr)
    run = paragraph.add_run(text)
    copy_run_properties(run, claim_rpr)
    return paragraph


def add_plain_run(paragraph, text: str) -> None:
    if not text:
        return
    run = paragraph.add_run(text)
    set_east_asia_font(run, "宋体")
    run.font.size = Pt(12)


def add_spec_item(
    doc: Document,
    item: SpecItem,
    registry: MathRegistry,
    symbols: tuple[str, ...],
):
    if item.kind == "title":
        paragraph = doc.add_paragraph(style="Title")
        paragraph.add_run(item.text)
        return paragraph
    if item.kind == "heading":
        paragraph = doc.add_paragraph(style="Heading 1")
        run = paragraph.add_run(item.text)
        run.style = doc.styles[REQUIRED_CHARACTER_STYLE]
        return paragraph
    if item.kind == "formula":
        paragraph = doc.add_paragraph(style="正文2")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.add_run(registry.add(item.text))
        punctuation = paragraph.add_run("。")
        set_east_asia_font(punctuation, "宋体")
        punctuation.font.size = Pt(12)
        return paragraph

    paragraph = doc.add_paragraph(style="正文2")
    for is_math, value in tokenize_inline_math(item.text, symbols):
        if is_math:
            paragraph.add_run(registry.add(value))
        else:
            add_plain_run(paragraph, value)
    return paragraph


def image_size(path: Path, max_width_in: float, max_height_in: float) -> tuple[float, float]:
    with Image.open(path) as image:
        width_px, height_px = image.size
    scale = min(max_width_in / width_px, max_height_in / height_px)
    return width_px * scale, height_px * scale


def add_figure(doc: Document, figure: FigureSpec, page_break_before: bool):
    width, height = image_size(figure.path, max_width_in=6.25, max_height_in=8.75)
    image_p = doc.add_paragraph(style="附图图号")
    image_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_p.paragraph_format.space_before = Pt(0)
    image_p.paragraph_format.space_after = Pt(0)
    image_p.paragraph_format.keep_with_next = True
    image_p.paragraph_format.page_break_before = page_break_before
    image_p.add_run().add_picture(str(figure.path), width=Inches(width), height=Inches(height))

    caption_p = doc.add_paragraph(style="附图图号")
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.paragraph_format.space_before = Pt(0)
    caption_p.paragraph_format.space_after = Pt(0)
    caption_p.paragraph_format.keep_together = True
    caption_run = caption_p.add_run(f"图{figure.number}")
    set_east_asia_font(caption_run, "宋体")
    caption_run.font.size = Pt(12)
    return caption_p


def add_abstract_figure(doc: Document, figure: FigureSpec):
    width, height = image_size(figure.path, max_width_in=6.10, max_height_in=8.85)
    paragraph = doc.add_paragraph(style="附图图号")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.add_run().add_picture(str(figure.path), width=Inches(width), height=Inches(height))
    return paragraph


def require_word_automation():
    if sys.platform != "win32":
        raise RuntimeError("原生 Word 公式转换仅支持安装 Microsoft Word 的 Windows 环境")
    try:
        import win32com.client as win32
    except ImportError as exc:
        raise RuntimeError("缺少 pywin32，无法调用 Microsoft Word 生成原生公式") from exc
    return win32


def process_with_word(path: Path, registry: MathRegistry, pdf_path: Path | None) -> int:
    win32 = require_word_automation()
    word = win32.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    document = None
    try:
        document = word.Documents.Open(str(path.resolve()), False, False)
        for spec in registry.specs:
            search = document.Content.Duplicate
            find = search.Find
            find.ClearFormatting()
            find.Text = spec.marker
            find.Forward = True
            find.Wrap = 0
            if not find.Execute():
                raise RuntimeError(f"未找到公式占位符：{spec.marker}")
            start = search.Start
            search.Text = spec.linear
            search.SetRange(start, start + len(spec.linear))
            search.OMaths.Add(search)
            search.OMaths.Item(1).BuildUp()
        for story in document.StoryRanges:
            with suppress(Exception):
                story.Fields.Update()
        document.Repaginate()
        page_count = int(document.ComputeStatistics(2))
        document.Save()
        if pdf_path is not None:
            pdf_path.parent.mkdir(parents=True, exist_ok=True)
            document.ExportAsFixedFormat(str(pdf_path.resolve()), 17)
        document.Close(False)
        document = None
        return page_count
    finally:
        if document is not None:
            document.Close(False)
        word.Quit()


def render_pdf_to_png(pdf_path: Path, render_dir: Path) -> list[Path]:
    executable = shutil.which("pdftoppm")
    if executable is None:
        raise RuntimeError("未找到 pdftoppm，无法把 Word 导出的 PDF 渲染为逐页 PNG")
    # Windows 版 Poppler 对包含中文的路径兼容性不稳定。使用系统临时目录中的
    # ASCII 文件名完成栅格化，再复制回案件工作区，避免路径编码导致假失败。
    with tempfile.TemporaryDirectory(prefix="cn_patent_docx_render_") as staging_raw:
        staging = Path(staging_raw)
        staged_pdf = staging / "input.pdf"
        shutil.copy2(pdf_path, staged_pdf)
        prefix = staging / "page"
        subprocess.run(
            [executable, "-png", "-r", "120", str(staged_pdf), str(prefix)],
            check=True,
        )
        staged_pages = sorted(staging.glob("page-*.png"))
        if not staged_pages:
            raise RuntimeError("pdftoppm 未生成逐页 PNG")
        pages: list[Path] = []
        for source in staged_pages:
            target = render_dir / source.name
            shutil.copy2(source, target)
            pages.append(target)
        return pages


def inspect_package(
    path: Path,
    expected_math_count: int,
    heading_style_id: str,
    strong_style_id: str,
) -> dict:
    with ZipFile(path) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    namespaces = {
        "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    }
    math_count = len(root.xpath(".//m:oMath", namespaces=namespaces))
    strong_count = len(
        root.xpath(
            './/w:p[w:pPr/w:pStyle[@w:val=$heading]]/w:r/w:rPr/w:rStyle[@w:val=$strong]',
            namespaces=namespaces,
            heading=heading_style_id,
            strong=strong_style_id,
        )
    )
    markers = root.xpath('.//w:t[contains(text(), "[[EQ")]/text()', namespaces=namespaces)
    if markers:
        raise RuntimeError(f"仍有公式占位符未转换：{markers}")
    if math_count != expected_math_count:
        raise RuntimeError(f"公式对象数量错误：期望 {expected_math_count}，实际 {math_count}")
    return {"math_count": math_count, "strong_heading_run_count": strong_count}


def build_document(
    source_dir: Path,
    template_path: Path,
    output_path: Path,
    render_dir: Path | None,
) -> dict:
    claims_path = source_dir / "权利要求书.md"
    spec_path = source_dir / "说明书.md"
    abstract_path = source_dir / "说明书摘要.md"
    figure_index_path = source_dir / "说明书附图.md"
    for path in [template_path, claims_path, spec_path, abstract_path, figure_index_path]:
        if not path.is_file():
            raise FileNotFoundError(path)

    claims = parse_claims(claims_path)
    specification = parse_specification(spec_path)
    abstract, abstract_figure_number = parse_abstract(abstract_path)
    figures = parse_figures(figure_index_path)
    by_number = {figure.number: figure for figure in figures}
    if abstract_figure_number not in by_number:
        raise ValueError(f"摘要附图图号不存在：图{abstract_figure_number}")
    symbols = collect_math_symbols(specification)
    registry = MathRegistry()

    doc = Document(template_path)
    section_props = validate_template(doc)
    heading_style_id = doc.styles["Heading 1"].style_id
    strong_style_id = doc.styles[REQUIRED_CHARACTER_STYLE].style_id
    claim_ppr = deepcopy(doc.paragraphs[0]._p.pPr)
    claim_rpr = next(
        (deepcopy(run._r.rPr) for run in doc.paragraphs[0].runs if run.text),
        None,
    )
    clear_body_keep_final_sectpr(doc, section_props[-1])

    last = None
    for claim in claims:
        last = add_claim(doc, claim, claim_ppr, claim_rpr)
    if last is None:
        raise ValueError("权利要求为空")
    end_section(last, section_props[0])

    for item in specification:
        last = add_spec_item(doc, item, registry, symbols)
    end_section(last, section_props[1])

    for index, figure in enumerate(figures):
        last = add_figure(doc, figure, page_break_before=index > 0)
    end_section(last, section_props[2])

    abstract_p = doc.add_paragraph(style="正文2")
    abstract_run = abstract_p.add_run(abstract)
    set_east_asia_font(abstract_run, "宋体")
    abstract_run.font.size = Pt(12)
    end_section(abstract_p, section_props[3])
    add_abstract_figure(doc, by_number[abstract_figure_number])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    title = next(item.text for item in specification if item.kind == "title")
    doc.core_properties.title = title
    doc.core_properties.subject = "中国发明专利申请文件"
    doc.save(output_path)

    pdf_path = None
    if render_dir is not None:
        render_dir.mkdir(parents=True, exist_ok=True)
        pdf_path = render_dir / f"{output_path.stem}.pdf"
    page_count = process_with_word(output_path, registry, pdf_path)
    package = inspect_package(
        output_path, len(registry.specs), heading_style_id, strong_style_id
    )

    check = Document(output_path)
    headers = [
        section.header.paragraphs[0].text.strip() if section.header.paragraphs else ""
        for section in check.sections
    ]
    heading_texts = [p.text for p in check.paragraphs if p.style.name == "Heading 1"]
    if headers != EXPECTED_HEADERS:
        raise RuntimeError(f"输出页眉错误：{headers}")
    if package["strong_heading_run_count"] != len(heading_texts):
        raise RuntimeError(
            "说明书小标题未全部应用“要点”样式："
            f"标题 {len(heading_texts)}，要点样式 {package['strong_heading_run_count']}"
        )
    if len(check.inline_shapes) != len(figures) + 1:
        raise RuntimeError("说明书附图或摘要附图数量错误")

    rendered_pages: list[Path] = []
    if pdf_path is not None:
        rendered_pages = render_pdf_to_png(pdf_path, render_dir)
        if len(rendered_pages) != page_count:
            raise RuntimeError(
                f"Word 页数与 PNG 页数不一致：Word {page_count}，PNG {len(rendered_pages)}"
            )

    return {
        "schema": REPORT_SCHEMA,
        "status": (
            "STRUCTURE_VERIFIED_VISUAL_REVIEW_PENDING"
            if render_dir is not None
            else "STRUCTURE_VERIFIED"
        ),
        "output": str(output_path.resolve()),
        "output_sha256": file_sha256(output_path),
        "inputs": {
            "template": str(template_path.resolve()),
            "template_sha256": file_sha256(template_path),
            "source_dir": str(source_dir.resolve()),
            "artifacts": [
                {"artifact_id": "claims", "path": str(claims_path.resolve()), "sha256": file_sha256(claims_path)},
                {"artifact_id": "specification", "path": str(spec_path.resolve()), "sha256": file_sha256(spec_path)},
                {"artifact_id": "abstract", "path": str(abstract_path.resolve()), "sha256": file_sha256(abstract_path)},
                {"artifact_id": "figure_index", "path": str(figure_index_path.resolve()), "sha256": file_sha256(figure_index_path)},
            ] + [
                {"artifact_id": f"figure_{figure.number}", "path": str(figure.path.resolve()), "sha256": file_sha256(figure.path)}
                for figure in figures
            ],
        },
        "evidence_scope": {
            "proves": [
                "DOCX ZIP/XML、分节、页眉、样式、原生公式对象和嵌入图片数量满足本脚本合同",
                "输出 DOCX 绑定当前模板、四文书源文件和每幅嵌入图片的 SHA-256",
            ],
            "does_not_prove": [
                "申请文件的法律实体条件已经满足",
                "未请求视觉检查时的逐页视觉效果",
                "附图本身不存在视觉缺陷",
            ],
        },
        "counts": {
            "claims": len(claims),
            "specification_items": len(specification),
            "headings_with_strong_style": package["strong_heading_run_count"],
            "native_word_math_objects": package["math_count"],
            "figures": len(figures),
            "inline_images": len(check.inline_shapes),
            "sections": len(check.sections),
            "pages": page_count,
        },
        "headers": headers,
        "abstract_figure": abstract_figure_number,
        "render": {
            "requested": render_dir is not None,
            "pdf": str(pdf_path.resolve()) if pdf_path else None,
            "png_pages": [str(path.resolve()) for path in rendered_pages],
            "visual_review_completed": False if render_dir is not None else None,
        },
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--case-dir", required=True, type=Path, help="案件根目录")
    parser.add_argument("--template", type=Path, help="输出模板，默认 <case-dir>/输出模版.docx")
    parser.add_argument("--source-dir", type=Path, help="申请文件目录，默认 <case-dir>/02-申请文件")
    parser.add_argument("--output", type=Path, help="输出 DOCX 路径")
    parser.add_argument(
        "--work-dir",
        type=Path,
        help="报告目录；仅在视觉检查时同时保存 PDF/PNG",
    )
    parser.add_argument(
        "--visual-review",
        action="store_true",
        help="仅在用户明确要求视觉检查时使用：导出 PDF 并生成逐页 PNG",
    )
    parser.add_argument(
        "--no-render",
        action="store_true",
        help=argparse.SUPPRESS,  # 向后兼容；默认已不渲染
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    case_dir = args.case_dir.resolve()
    template = (args.template or case_dir / "输出模版.docx").resolve()
    source_dir = (args.source_dir or case_dir / "02-申请文件").resolve()
    title_items = parse_specification(source_dir / "说明书.md")
    title = next(item.text for item in title_items if item.kind == "title")
    safe_title = re.sub(r"[^\w\-\u4e00-\u9fff]", "_", title).strip("_")
    output = (args.output or case_dir / f"{safe_title}-专利申请文件.docx").resolve()
    work_dir = (
        args.work_dir
        or case_dir / "03-审查工作区" / f"docx组装-{datetime.now():%Y%m%d-%H%M%S}"
    ).resolve()
    if args.visual_review and args.no_render:
        raise ValueError("--visual-review 与兼容参数 --no-render 不能同时使用")
    render_dir = work_dir / "render" if args.visual_review else None

    report = build_document(source_dir, template, output, render_dir)
    work_dir.mkdir(parents=True, exist_ok=True)
    report_path = work_dir / "docx-assembly-report.json"
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    summary = {
        "output": str(output),
        "report": str(report_path),
        "status": report["status"],
        "visual_review_requested": report["render"]["requested"],
        **report["counts"],
    }
    print(json.dumps(summary, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
