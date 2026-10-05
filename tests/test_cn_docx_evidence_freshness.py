"""DOCX 组装证据的新鲜度回归，不依赖 Word。"""
from __future__ import annotations

import hashlib
import importlib.util
import json
import sys
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "skills/cn-patent-application-creator/scripts/verify_docx_assembly.py"


def load_module():
    spec = importlib.util.spec_from_file_location("verify_docx_assembly", SCRIPT)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()



def write_minimal_docx(path: Path, *, missing_ignorable_namespace: bool = False) -> None:
    declarations = (
        'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml"'
    )
    if not missing_ignorable_namespace:
        declarations += ' xmlns:w15="http://schemas.microsoft.com/office/word/2012/wordml"'
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:document {declarations} mc:Ignorable="w14 w15"><w:body><w:p/></w:body></w:document>'
    ).encode("utf-8")
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", b"<Types xmlns=\"http://schemas.openxmlformats.org/package/2006/content-types\"/>")
        archive.writestr("_rels/.rels", b"<Relationships xmlns=\"http://schemas.openxmlformats.org/package/2006/relationships\"/>")
        archive.writestr("word/document.xml", document)
        archive.writestr("word/_rels/document.xml.rels", b"<Relationships xmlns=\"http://schemas.openxmlformats.org/package/2006/relationships\"/>")

def _write_embedded_figure_docx(path: Path, image: Path) -> None:
    """构造带说明书附图和摘要附图的最小 DOCX，摘要复用同一幅图的字节。"""
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.style import WD_STYLE_TYPE

    document = Document()
    for name in ("Normal (Web)", "正文2", "附图图号"):
        if name not in document.styles:
            document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    headers = ["权利要求书", "说明书", "说明书附图", "说明书摘要", "摘要附图"]
    document.sections[0].header.paragraphs[0].text = headers[0]
    document.add_paragraph("权利要求")
    for header in headers[1:]:
        section = document.add_section(WD_SECTION.NEW_PAGE)
        section.header.is_linked_to_previous = False
        section.header.paragraphs[0].text = header
        if header == "说明书附图":
            document.add_paragraph().add_run().add_picture(str(image))
            document.add_paragraph("图1")
        elif header == "摘要附图":
            document.add_paragraph().add_run().add_picture(str(image))
        else:
            document.add_paragraph(header)
    document.save(path)


def test_docx_report_binds_every_source_and_detects_stale_input(tmp_path):
    module = load_module()
    template = tmp_path / "template.docx"
    output = tmp_path / "output.docx"
    artifacts = []
    for artifact_id, name in [("claims", "权利要求书.md"), ("specification", "说明书.md"), ("abstract", "说明书摘要.md"), ("figure_index", "说明书附图.md"), ("figure_1", "图1.png")]:
        path = tmp_path / name
        if artifact_id == "figure_1":
            from PIL import Image, ImageDraw
            image = Image.new("RGB", (80, 60), "white")
            ImageDraw.Draw(image).rectangle((8, 8, 71, 51), outline="black", width=2)
            image.save(path)
        elif artifact_id == "abstract":
            path.write_text("摘要附图：图1。\n", encoding="utf-8")
        else:
            path.write_bytes(artifact_id.encode("utf-8"))
        artifacts.append({"artifact_id": artifact_id, "path": str(path), "sha256": digest(path)})
    template.write_bytes(b"template")
    _write_embedded_figure_docx(output, tmp_path / "图1.png")
    report_path = tmp_path / "docx-assembly-report.json"
    report_path.write_text(json.dumps({
        "schema": "cn-patent-docx-assembly/v2", "status": "STRUCTURE_VERIFIED",
        "output": str(output), "output_sha256": digest(output),
        "abstract_figure": 1,
        "inputs": {"template": str(template), "template_sha256": digest(template), "source_dir": str(tmp_path), "artifacts": artifacts},
    }, ensure_ascii=False), encoding="utf-8")
    verified = module.verify(report_path)
    assert verified["status"] == "PASS", verified
    (tmp_path / "说明书.md").write_text("changed", encoding="utf-8")
    stale = module.verify(report_path)
    assert stale["status"] == "FAIL"
    assert any(item["code"] == "DOCX-EVIDENCE-STALE" for item in stale["errors"])


def test_docx_report_rejects_undeclared_mc_ignorable_namespace(tmp_path):
    module = load_module()
    template = tmp_path / "template.docx"
    output = tmp_path / "output.docx"
    template.write_bytes(b"template")
    write_minimal_docx(output, missing_ignorable_namespace=True)
    artifacts = []
    for artifact_id in ("claims", "specification", "abstract", "figure_index"):
        source = tmp_path / f"{artifact_id}.txt"
        source.write_text(artifact_id, encoding="utf-8")
        artifacts.append({"artifact_id": artifact_id, "path": str(source), "sha256": digest(source)})
    report_path = tmp_path / "report.json"
    report_path.write_text(json.dumps({
        "schema": "cn-patent-docx-assembly/v2",
        "status": "STRUCTURE_VERIFIED",
        "output": str(output),
        "output_sha256": digest(output),
        "inputs": {"template": str(template), "template_sha256": digest(template), "artifacts": artifacts},
    }), encoding="utf-8")
    result = module.verify(report_path)
    assert result["status"] == "FAIL"
    assert any(item["code"] == "DOCX-PACKAGE-IGNORABLE-NS" for item in result["errors"])


def test_invalid_zip_docx_returns_structured_failure(tmp_path):
    """损坏 ZIP 必须返回稳定失败字段，不能把 BadZipFile 抛成未捕获崩溃。"""
    module = load_module()
    template = tmp_path / "template.docx"
    output = tmp_path / "output.docx"
    template.write_bytes(b"template")
    output.write_bytes(b"this-is-not-a-zip")
    artifacts = []
    for artifact_id in ("claims", "specification", "abstract", "figure_index"):
        source = tmp_path / f"{artifact_id}.txt"
        source.write_text(artifact_id, encoding="utf-8")
        artifacts.append({"artifact_id": artifact_id, "path": str(source), "sha256": digest(source)})
    report_path = tmp_path / "report.json"
    report_path.write_text(json.dumps({
        "schema": "cn-patent-docx-assembly/v2",
        "status": "STRUCTURE_VERIFIED",
        "copy_kind": "review",
        "pending_marks": [],
        "output": str(output),
        "output_sha256": digest(output),
        "inputs": {"template": str(template), "template_sha256": digest(template), "artifacts": artifacts},
    }), encoding="utf-8")
    result = module.verify(report_path)
    assert result["status"] == "FAIL"
    assert result["schema_id"] == "cn-patent-docx-assembly-verification/v1"
    assert "errors" in result and "evidence_scope" in result
    assert any(item["code"] == "DOCX-PACKAGE-ZIP" for item in result["errors"])
