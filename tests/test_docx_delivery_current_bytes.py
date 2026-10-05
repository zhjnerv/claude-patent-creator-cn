"""F02 专项：DOCX 交付必须独立复算当前 DOCX/组装报告字节，不得信任旧 freshness。

覆盖：
  * 合法新组装 PASS（端到端：assemble_application_docx → verify_docx_assembly →
    verify_drawing_docx_delivery）。
  * 修改、替换、删除真实 DOCX 内嵌图片/正文后，delivery verifier 失败。
  * drawing verification 必须以 figures[] 数组绑定当前 PNG/Draw.io 字节，
    旧 fixture（仅 brief_sha256）必须被拒。
"""
from __future__ import annotations

import hashlib
import importlib.util
import json
import re
import subprocess
import sys
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

ROOT = Path(__file__).resolve().parents[1]

ASSEMBLER_PATH = (
    ROOT / "skills/cn-patent-application-creator/scripts/assemble_application_docx.py"
)
VERIFY_ASSEMBLY_PATH = (
    ROOT / "skills/cn-patent-application-creator/scripts/verify_docx_assembly.py"
)
DELIVERY_PATH = (
    ROOT / "skills/cn-patent-diagram-generator/scripts/verify_drawing_docx_delivery.py"
)


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _write_json(path: Path, value: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )


def _load_module(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    assert spec is not None and spec.loader is not None, path
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def _write_test_template(path: Path) -> None:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    document = Document()
    for name in ("Normal (Web)", "正文2", "附图图号"):
        if name not in document.styles:
            document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    paragraph = document.add_paragraph()
    paragraph.add_run("模板权利要求")
    ppr = paragraph._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), "1")
    numpr.extend((ilvl, numid))
    ppr.append(numpr)
    headers = ["权利要求书", "说明书", "说明书附图", "说明书摘要", "摘要附图"]
    document.sections[0].header.paragraphs[0].text = headers[0]
    for header in headers[1:]:
        section = document.add_section(WD_SECTION.NEW_PAGE)
        section.header.is_linked_to_previous = False
        section.header.paragraphs[0].text = header
    document.save(path)


def _write_source(source_dir: Path) -> Path:
    source_dir.mkdir(parents=True, exist_ok=True)
    (source_dir / "权利要求书.md").write_text(
        "# 权利要求书\n\n1. 一种方法，其特征在于，包括：S101：执行第一步骤；S102：执行第二步骤。\n",
        encoding="utf-8",
    )
    (source_dir / "说明书.md").write_text(
        "# 示例方法\n\n"
        "## 技术领域\n涉及示例技术。\n\n"
        "## 背景技术\n现有方案不足。\n\n"
        "## 发明内容\n提供示例方案。\n\n"
        "## 附图说明\n图1为本申请一实施方式中的流程图。\n\n"
        "图中：S101-第一步骤；S102-第二步骤。\n\n"
        "## 具体实施方式\n以下结合附图说明本发明的具体实施例。\n\n"
        "### 实施例1\n如图1所示，先执行步骤S101，再执行步骤S102。\n",
        encoding="utf-8",
    )
    (source_dir / "说明书摘要.md").write_text(
        "# 示例方法\n\n本发明提供一种示例方法。\n\n摘要附图：图1。\n",
        encoding="utf-8",
    )
    (source_dir / "说明书附图.md").write_text(
        "# 说明书附图\n\n## 图1 流程图\n\n![图1](说明书附图/图1.png)\n",
        encoding="utf-8",
    )
    return source_dir / "说明书附图/图1.png"


def _make_png(path: Path) -> None:
    from PIL import Image, ImageDraw

    path.parent.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (200, 100), "white")
    ImageDraw.Draw(image).rectangle((10, 10, 189, 89), outline="black", width=2)
    image.save(path)


def _assemble(source_dir: Path, template: Path, docx: Path) -> dict:
    assembler = _load_module("cn_docx_assembler_for_f02", ASSEMBLER_PATH)
    return assembler.build_document(source_dir, template, docx, None)


def _verify_assembly(report_path: Path) -> dict:
    cmd = [
        sys.executable,
        str(VERIFY_ASSEMBLY_PATH),
        "--report",
        str(report_path),
    ]
    result = subprocess.run(cmd, capture_output=True, text=True, check=False, cwd=ROOT)
    return {
        "returncode": result.returncode,
        "stdout": result.stdout,
        "stderr": result.stderr,
    }


def _build_drawing_files(
    case_dir: Path,
    brief_path: Path,
    png_path: Path,
    drawio_path: Path,
    *,
    include_figures: bool = True,
) -> Path:
    figures_block = []
    if include_figures:
        figures_block.append({
            "figure_number": 1,
            "png": {"path": str(png_path), "sha256": _sha(png_path)},
            "drawio": {"path": str(drawio_path), "sha256": _sha(drawio_path)},
        })
    drawing_path = case_dir / "drawing-verification.json"
    _write_json(drawing_path, {
        "status": "PASS",
        "errors": [],
        "brief_sha256": _sha(brief_path),
        "figures": figures_block,
    })
    return drawing_path


def _run_delivery(
    case_dir: Path,
    brief_path: Path,
    drawing_path: Path,
    docx_report_path: Path,
    docx_verification_path: Path,
):
    output_path = case_dir / "delivery.json"
    cmd = [
        sys.executable,
        str(DELIVERY_PATH),
        "--case-dir",
        str(case_dir),
        "--brief",
        str(brief_path),
        "--drawing-verification",
        str(drawing_path),
        "--docx-report",
        str(docx_report_path),
        "--docx-verification",
        str(docx_verification_path),
        "--output",
        str(output_path),
    ]
    result = subprocess.run(cmd, capture_output=True, text=True, check=False, cwd=ROOT)
    payload = json.loads(output_path.read_text(encoding="utf-8"))
    return result, payload


def _build_realistic_docx_evidence(tmp_path: Path) -> dict:
    """端到端构造合法组装→合法 verification→合法交付证据。"""
    case_dir = tmp_path / "case"
    case_dir.mkdir(parents=True, exist_ok=True)
    source = case_dir / "02-申请文件"
    template = case_dir / "template.docx"
    docx = case_dir / "output.docx"
    _write_test_template(template)
    png_path = _write_source(source)
    _make_png(png_path)
    drawio_path = case_dir / "drawings/图1.drawio"
    drawio_path.parent.mkdir(parents=True, exist_ok=True)
    drawio_path.write_bytes(b"<mxfile>valid-drawio</mxfile>")

    report = _assemble(source, template, docx)
    report_path = case_dir / "docx-report.json"
    _write_json(report_path, report)
    assert report["status"] in {"STRUCTURE_VERIFIED", "STRUCTURE_VERIFIED_VISUAL_REVIEW_PENDING"}

    freshness = _verify_assembly(report_path)
    freshness_path = case_dir / "docx-verification.json"
    _write_json(freshness_path, json.loads(freshness["stdout"]))
    assert freshness["returncode"] == 0, freshness

    brief = case_dir / "brief.json"
    _write_json(brief, {
        "figures": [{
            "figure_number": 1,
            "outputs": {
                "final_png": str(png_path.relative_to(case_dir)),
                "drawio": str(drawio_path.relative_to(case_dir)),
            },
        }],
    })
    drawing = _build_drawing_files(case_dir, brief, png_path, drawio_path)
    return {
        "case_dir": case_dir,
        "docx": docx,
        "docx_report": report_path,
        "docx_verification": freshness_path,
        "brief": brief,
        "drawing": drawing,
        "png": png_path,
        "drawio": drawio_path,
    }


# ---------------------------------------------------------------------------
# 正例：合法新组装必须 PASS。
# ---------------------------------------------------------------------------

def test_legitimate_new_assembly_passes_delivery(tmp_path):
    env = _build_realistic_docx_evidence(tmp_path)
    result, payload = _run_delivery(
        env["case_dir"], env["brief"], env["drawing"],
        env["docx_report"], env["docx_verification"],
    )
    assert result.returncode == 0, result.stdout + result.stderr + json.dumps(payload, ensure_ascii=False)
    assert payload["status"] == "PASS"
    assert payload["figure_count"] == 1
    assert payload["docx_assembly_reverified_status"] == "PASS"


# ---------------------------------------------------------------------------
# 反例：修改真实 DOCX 内嵌图片/正文字节后必须 FAIL。
# ---------------------------------------------------------------------------

def test_modified_real_docx_embedded_png_fails_delivery(tmp_path):
    env = _build_realistic_docx_evidence(tmp_path)
    docx = env["docx"]
    with ZipFile(docx) as z:
        stored = [(info, z.read(info.filename)) for info in z.infolist()]
    with ZipFile(docx, "w", compression=ZIP_DEFLATED) as z:
        replacement = b"FAKE-EMBEDDED-PNG-BYTES"
        for info, value in stored:
            if info.filename.startswith("word/media/"):
                z.writestr(info, replacement)
            else:
                z.writestr(info, value)
    result, payload = _run_delivery(
        env["case_dir"], env["brief"], env["drawing"],
        env["docx_report"], env["docx_verification"],
    )
    assert result.returncode == 2, result.stdout + result.stderr
    codes = {item["code"] for item in payload["errors"]}
    assert "DELIVERY-DOCX-OUTPUT-STALE" in codes or "DELIVERY-DOCX-STALE" in codes


def test_modified_real_docx_document_xml_fails_delivery(tmp_path):
    """在 word/document.xml 末尾注入任意字节，必须被识别为当前 DOCX 已变。"""
    env = _build_realistic_docx_evidence(tmp_path)
    docx = env["docx"]
    with ZipFile(docx) as z:
        stored = [(info, z.read(info.filename)) for info in z.infolist()]
    with ZipFile(docx, "w", compression=ZIP_DEFLATED) as z:
        for info, value in stored:
            if info.filename == "word/document.xml":
                z.writestr(info, value + b"\n<!-- tampered -->\n")
            else:
                z.writestr(info, value)
    result, payload = _run_delivery(
        env["case_dir"], env["brief"], env["drawing"],
        env["docx_report"], env["docx_verification"],
    )
    assert result.returncode == 2, result.stdout + result.stderr
    codes = {item["code"] for item in payload["errors"]}
    assert codes & {"DELIVERY-DOCX-OUTPUT-STALE", "DELIVERY-DOCX-STALE"}


def test_replaced_real_docx_fails_delivery(tmp_path):
    """完全替换 DOCX 文件（同一路径不同字节）必须被捕获。"""
    env = _build_realistic_docx_evidence(tmp_path)
    env["docx"].write_bytes(b"totally-different-docx-bytes")
    result, payload = _run_delivery(
        env["case_dir"], env["brief"], env["drawing"],
        env["docx_report"], env["docx_verification"],
    )
    assert result.returncode == 2, result.stdout + result.stderr
    codes = {item["code"] for item in payload["errors"]}
    assert "DELIVERY-DOCX-OUTPUT-STALE" in codes or "DELIVERY-DOCX-STALE" in codes


def test_deleted_real_docx_fails_delivery(tmp_path):
    env = _build_realistic_docx_evidence(tmp_path)
    env["docx"].unlink()
    result, payload = _run_delivery(
        env["case_dir"], env["brief"], env["drawing"],
        env["docx_report"], env["docx_verification"],
    )
    assert result.returncode == 2, result.stdout + result.stderr
    codes = {item["code"] for item in payload["errors"]}
    assert "DELIVERY-DOCX-OUTPUT-STALE" in codes or "DELIVERY-DOCX-STALE" in codes


# ---------------------------------------------------------------------------
# 反例：附图验证必须绑定当前 PNG/Draw.io 字节。
# ---------------------------------------------------------------------------

def test_old_drawing_verification_without_figures_array_fails(tmp_path):
    env = _build_realistic_docx_evidence(tmp_path)
    stale_drawing = env["case_dir"] / "drawing-verification.json"
    _write_json(stale_drawing, {
        "status": "PASS",
        "errors": [],
        "brief_sha256": _sha(env["brief"]),
    })
    result, payload = _run_delivery(
        env["case_dir"], env["brief"], stale_drawing,
        env["docx_report"], env["docx_verification"],
    )
    assert result.returncode == 2, result.stdout + result.stderr
    codes = {item["code"] for item in payload["errors"]}
    assert "DELIVERY-DRAWING-BINDING" in codes


def test_changed_real_png_fails_drawing_binding(tmp_path):
    env = _build_realistic_docx_evidence(tmp_path)
    # 在装配之后修改真实 PNG 字节：drawing 报告里的 sha256 已不匹配当前文件。
    env["png"].write_bytes(b"modified-after-delivery")
    result, payload = _run_delivery(
        env["case_dir"], env["brief"], env["drawing"],
        env["docx_report"], env["docx_verification"],
    )
    assert result.returncode == 2, result.stdout + result.stderr
    codes = {item["code"] for item in payload["errors"]}
    assert "DELIVERY-DRAWING-EVIDENCE-STALE" in codes or "DELIVERY-DOCX-FIGURE-STALE" in codes


def test_changed_real_drawio_fails_drawing_binding(tmp_path):
    env = _build_realistic_docx_evidence(tmp_path)
    env["drawio"].write_bytes(b"<mxfile>modified</mxfile>")
    result, payload = _run_delivery(
        env["case_dir"], env["brief"], env["drawing"],
        env["docx_report"], env["docx_verification"],
    )
    assert result.returncode == 2, result.stdout + result.stderr
    codes = {item["code"] for item in payload["errors"]}
    assert "DELIVERY-DRAWING-EVIDENCE-STALE" in codes


# ---------------------------------------------------------------------------
# 旧 freshness 报告不能掩盖当前复算失败。
# ---------------------------------------------------------------------------

def test_stale_docx_verification_with_real_docx_mismatch_fails(tmp_path):
    env = _build_realistic_docx_evidence(tmp_path)
    # 篡改真实 DOCX，但保留 docx-verification.json（PASS）作为旧 freshness。
    env["docx"].write_bytes(b"tampered-after-freshness")
    result, payload = _run_delivery(
        env["case_dir"], env["brief"], env["drawing"],
        env["docx_report"], env["docx_verification"],
    )
    assert result.returncode == 2, result.stdout + result.stderr
    assert payload["docx_assembly_reverified_status"] == "FAIL"
    # 旧 freshness 不再被当作新鲜度证据；当前独立复算失败即可拒绝。
    codes = {item["code"] for item in payload["errors"]}
    assert "DELIVERY-DOCX-OUTPUT-STALE" in codes or "DELIVERY-DOCX-STALE" in codes


def _rewrite_docx(docx: Path, transform) -> None:
    with ZipFile(docx) as archive:
        stored = [(info, archive.read(info.filename)) for info in archive.infolist()]
    with ZipFile(docx, "w", compression=ZIP_DEFLATED) as archive:
        for info, value in stored:
            archive.writestr(info, transform(info.filename, value))


def _rebind_output_hash(report_path: Path, docx: Path) -> None:
    report = json.loads(report_path.read_text(encoding="utf-8"))
    report["output_sha256"] = _sha(docx)
    _write_json(report_path, report)


def _different_png(path: Path) -> bytes:
    from PIL import Image, ImageDraw

    image = Image.new("RGB", (90, 50), "white")
    ImageDraw.Draw(image).ellipse((5, 5, 84, 44), outline="red", width=4)
    image.save(path)
    return path.read_bytes()


def test_rebound_hash_still_rejects_replaced_embedded_png(tmp_path):
    """只更新整包哈希、替换内嵌 PNG 时，交付门禁仍必须失败。"""
    env = _build_realistic_docx_evidence(tmp_path)
    wrong = _different_png(tmp_path / "wrong.png")

    def transform(name, value):
        if name.startswith("word/media/") and name.endswith(".png"):
            return wrong
        return value

    _rewrite_docx(env["docx"], transform)
    _rebind_output_hash(env["docx_report"], env["docx"])
    result, payload = _run_delivery(
        env["case_dir"], env["brief"], env["drawing"], env["docx_report"], env["docx_verification"],
    )
    assert result.returncode == 2, result.stdout + result.stderr
    codes = {item["code"] for item in payload["errors"]}
    assert "DELIVERY-DOCX-EMBEDDED" in codes
    assert payload["status"] == "FAIL"


def _build_two_figure_evidence(tmp_path: Path) -> dict:
    from PIL import Image, ImageDraw

    case_dir = tmp_path / "case"
    case_dir.mkdir(parents=True, exist_ok=True)
    source = case_dir / "02-申请文件"
    source.mkdir(parents=True, exist_ok=True)
    (source / "权利要求书.md").write_text(
        "# 权利要求书\n\n1. 一种方法，其特征在于，包括：S101：执行第一步骤；S102：执行第二步骤。\n",
        encoding="utf-8",
    )
    (source / "说明书.md").write_text(
        "# 示例方法\n\n## 技术领域\n涉及示例技术。\n\n## 背景技术\n现有方案不足。\n\n"
        "## 发明内容\n提供示例方案。\n\n## 附图说明\n图1为本申请一实施方式中的流程图。\n\n"
        "图2为本申请一实施方式中的结构图。\n\n图中：S101-第一步骤；S102-第二步骤。\n\n"
        "## 具体实施方式\n以下结合附图说明本发明的具体实施例。\n\n### 实施例1\n"
        "如图1所示，先执行步骤S101。如图2所示，再执行步骤S102。\n",
        encoding="utf-8",
    )
    (source / "说明书摘要.md").write_text(
        "# 示例方法\n\n本发明提供一种示例方法。\n\n摘要附图：图1。\n",
        encoding="utf-8",
    )
    (source / "说明书附图.md").write_text(
        "# 说明书附图\n\n## 图1 流程图\n\n![图1](说明书附图/图1.png)\n\n## 图2 结构图\n\n![图2](说明书附图/图2.png)\n",
        encoding="utf-8",
    )
    figure_dir = source / "说明书附图"
    figure_dir.mkdir(parents=True, exist_ok=True)
    paths = []
    for number, color in ((1, "black"), (2, "blue")):
        image_path = figure_dir / f"图{number}.png"
        image = Image.new("RGB", (160, 80), "white")
        ImageDraw.Draw(image).rectangle((12, 12, 147, 67), outline=color, width=3)
        image.save(image_path)
        paths.append(image_path)
    drawio_paths = []
    for number in (1, 2):
        drawio_path = case_dir / f"drawings/图{number}.drawio"
        drawio_path.parent.mkdir(parents=True, exist_ok=True)
        drawio_path.write_bytes(f"<mxfile>figure-{number}</mxfile>".encode("utf-8"))
        drawio_paths.append(drawio_path)
    template = case_dir / "template.docx"
    docx = case_dir / "output.docx"
    _write_test_template(template)
    report = _assemble(source, template, docx)
    report_path = case_dir / "docx-report.json"
    _write_json(report_path, report)
    freshness = _verify_assembly(report_path)
    freshness_path = case_dir / "docx-verification.json"
    _write_json(freshness_path, json.loads(freshness["stdout"]))
    assert freshness["returncode"] == 0, freshness
    brief = case_dir / "brief.json"
    _write_json(brief, {
        "figures": [
            {
                "figure_number": number,
                "outputs": {
                    "final_png": str(image_path.relative_to(case_dir)),
                    "drawio": str(drawio_path.relative_to(case_dir)),
                },
            }
            for number, image_path, drawio_path in zip((1, 2), paths, drawio_paths)
        ],
    })
    drawing = case_dir / "drawing-verification.json"
    _write_json(drawing, {
        "status": "PASS",
        "errors": [],
        "brief_sha256": _sha(brief),
        "figures": [
            {
                "figure_number": number,
                "png": {"path": str(image_path), "sha256": _sha(image_path)},
                "drawio": {"path": str(drawio_path), "sha256": _sha(drawio_path)},
            }
            for number, image_path, drawio_path in zip((1, 2), paths, drawio_paths)
        ],
    })
    return {
        "case_dir": case_dir,
        "docx": docx,
        "docx_report": report_path,
        "docx_verification": freshness_path,
        "brief": brief,
        "drawing": drawing,
    }


def test_same_figure_reused_by_abstract_passes(tmp_path):
    env = _build_two_figure_evidence(tmp_path)
    with ZipFile(env["docx"]) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
    references = xml.count("r:embed=")
    assert references == 3
    assert len(media) < references
    result, payload = _run_delivery(
        env["case_dir"], env["brief"], env["drawing"], env["docx_report"], env["docx_verification"],
    )
    assert result.returncode == 0, result.stdout + result.stderr + json.dumps(payload, ensure_ascii=False)
    assert payload["status"] == "PASS"
    assert payload["figure_count"] == 2


def test_swapped_same_count_images_fail_after_rebind(tmp_path):
    env = _build_two_figure_evidence(tmp_path)
    with ZipFile(env["docx"]) as archive:
        media = [(info.filename, archive.read(info.filename)) for info in archive.infolist() if info.filename.startswith("word/media/") and info.filename.endswith(".png")]
    assert len(media) == 2
    swapped = {media[0][0]: media[1][1], media[1][0]: media[0][1]}

    def rewrite(name, value):
        return swapped.get(name, value)

    _rewrite_docx(env["docx"], rewrite)
    _rebind_output_hash(env["docx_report"], env["docx"])
    result, payload = _run_delivery(
        env["case_dir"], env["brief"], env["drawing"], env["docx_report"], env["docx_verification"],
    )
    assert result.returncode == 2, result.stdout + result.stderr
    assert any(item["code"] == "DELIVERY-DOCX-EMBEDDED" for item in payload["errors"])


def test_deleted_image_reference_fails_after_rebind(tmp_path):
    env = _build_two_figure_evidence(tmp_path)

    def rewrite(name, value):
        if name != "word/document.xml":
            return value
        text = value.decode("utf-8")
        start = text.find("<a:blip")
        end = text.find("/>", start)
        assert start != -1 and end != -1
        return (text[:start] + text[end + 2:]).encode("utf-8")

    _rewrite_docx(env["docx"], rewrite)
    _rebind_output_hash(env["docx_report"], env["docx"])
    result, payload = _run_delivery(
        env["case_dir"], env["brief"], env["drawing"], env["docx_report"], env["docx_verification"],
    )
    assert result.returncode == 2, result.stdout + result.stderr
    assert any(item["code"] == "DELIVERY-DOCX-EMBEDDED" for item in payload["errors"])


def test_wrong_abstract_figure_fails_after_rebind(tmp_path):
    env = _build_two_figure_evidence(tmp_path)

    def rewrite(name, value):
        if name != "word/document.xml":
            return value
        text = value.decode("utf-8")
        matches = list(re.finditer(r'r:embed="(rId\d+)"', text))
        assert len(matches) >= 3
        figure_two = matches[1].group(1)
        last = matches[-1]
        text = text[:last.start(1)] + figure_two + text[last.end(1):]
        return text.encode("utf-8")

    _rewrite_docx(env["docx"], rewrite)
    _rebind_output_hash(env["docx_report"], env["docx"])
    result, payload = _run_delivery(
        env["case_dir"], env["brief"], env["drawing"], env["docx_report"], env["docx_verification"],
    )
    assert result.returncode == 2, result.stdout + result.stderr
    messages = " ".join(item["message"] for item in payload["errors"])
    assert "摘要附图" in messages


def test_duplicate_figure_section_is_rejected_after_rebind(tmp_path):
    """不能只核对第一个同名分节而放过后续显示的图片。"""
    from copy import deepcopy
    from lxml import etree

    env = _build_realistic_docx_evidence(tmp_path)
    verifier = _load_module("duplicate_section_verifier", VERIFY_ASSEMBLY_PATH)

    def rewrite(name, value):
        if name != "word/document.xml":
            return value
        root = etree.fromstring(value)
        body = root.find(verifier.W + "body")
        blocks = verifier._section_blocks(body)[2][1]
        insertion = body.index(blocks[-1]) + 1
        for offset, block in enumerate(blocks):
            body.insert(insertion + offset, deepcopy(block))
        return etree.tostring(root, encoding="utf-8", xml_declaration=True, standalone=True)

    _rewrite_docx(env["docx"], rewrite)
    _rebind_output_hash(env["docx_report"], env["docx"])
    result, payload = _run_delivery(
        env["case_dir"], env["brief"], env["drawing"], env["docx_report"], env["docx_verification"],
    )
    assert result.returncode == 2, result.stdout + result.stderr
    assert any("各有一个" in item["message"] for item in payload["errors"])


def test_corrupt_package_does_not_claim_verified_images(tmp_path):
    env = _build_realistic_docx_evidence(tmp_path)
    env["docx"].write_bytes(b"not-a-docx")
    _rebind_output_hash(env["docx_report"], env["docx"])
    result = _verify_assembly(env["docx_report"])
    assert result["returncode"] == 2
    payload = json.loads(result["stdout"])
    assert payload["evidence_scope"]["proves"] == []
    assert any(item["code"] == "DOCX-PACKAGE-ZIP" for item in payload["errors"])
