"""中国发明专利申请 DOCX 组装合同回归测试。"""

from __future__ import annotations

import importlib.util
import shutil
import sys
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = (
    ROOT
    / "skills"
    / "cn-patent-application-creator"
    / "scripts"
    / "assemble_application_docx.py"
)
REFERENCE = (
    ROOT
    / "skills"
    / "cn-patent-application-creator"
    / "references"
    / "docx-assembly.md"
)


def _load_module():
    spec = importlib.util.spec_from_file_location("cn_docx_assembler", SCRIPT)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


ASSEMBLER = _load_module()


def _write_test_template(path: Path) -> None:
    document = Document()
    for name in ("Normal (Web)", "正文2", "附图图号"):
        if name not in document.styles:
            document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    paragraph = document.add_paragraph()
    paragraph.add_run("模板权利要求")
    ppr = paragraph._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), "1")
    numpr.extend((ilvl, numid))
    ppr.append(numpr)
    headers = ["权利要求书", "说明书", "说明书附图", "说明书摘要", "摘要附图"]
    document.sections[0].header.paragraphs[0].text = headers[0]
    for header in headers[1:]:
        section = document.add_section(WD_SECTION.NEW_PAGE)
        section.header.is_linked_to_previous = False
        section.header.paragraphs[0].text = header
    document.save(path)


def test_docx_assembly_skill_assets_are_utf8_and_documented():
    for path in (SCRIPT, REFERENCE):
        raw = path.read_bytes()
        assert raw
        assert not raw.startswith(b"\xef\xbb\xbf")
        raw.decode("utf-8")

    skill = (
        ROOT / "skills" / "cn-patent-application-creator" / "SKILL.md"
    ).read_text(encoding="utf-8")
    for required in (
        "assemble_application_docx.py",
        "references/docx-assembly.md",
        "m:oMath",
        "Strong",
        "visual_review_completed",
    ):
        assert required in skill
    assert "视觉检查默认关闭" in skill
    agents = (ROOT / "AGENTS.md").read_text(encoding="utf-8")
    readme = (ROOT / "README.md").read_text(encoding="utf-8")
    assert "只有用户明确要求视觉检查时才导出 PDF/PNG" in agents
    assert "不提供业务型 Slash Command" in readme
    assert "python-docx" in readme


def test_formula_parser_supports_explicit_and_legacy_blocks(tmp_path):
    module = ASSEMBLER
    spec = tmp_path / "说明书.md"
    spec.write_text(
        """# 示例装置

## 技术领域

患者年龄为 $a$，转换系数为 $C(P,a)$。

## 背景技术

$$
C(P,a)=C(P,a_i)+(a-a_i)×[C(P,a_(i+1))-C(P,a_i)]/(a_(i+1)-a_i)
$$

## 发明内容

```math
ED_CT=Q_CT×C_CT(P,a)
```

## 附图说明

图1是示意图。

## 具体实施方式

m_k=m_(k-1)+(x_k-m_(k-1))/k。
""",
        encoding="utf-8",
    )

    items = module.parse_specification(spec)
    assert [item.kind for item in items].count("formula") == 3
    assert [item.text for item in items if item.kind == "heading"] == [
        "技术领域",
        "背景技术",
        "发明内容",
        "附图说明",
        "具体实施方式",
    ]
    symbols = module.collect_math_symbols(items)
    tokens = module.tokenize_inline_math("患者年龄为 $a$，转换系数为 $C(P,a)$。", symbols)
    assert [value for is_math, value in tokens if is_math] == ["a", "C(P,a)"]


def test_claim_abstract_and_figure_contracts(tmp_path):
    module = ASSEMBLER
    claims = tmp_path / "权利要求书.md"
    claims.write_text("# 权利要求书\n\n1. 第一项。\n\n2. 第二项。\n", encoding="utf-8")
    assert module.parse_claims(claims) == ["第一项。", "第二项。"]

    compact_claims = tmp_path / "紧凑权利要求书.md"
    compact_claims.write_text(
        "# 权利要求书\n\n1. 第一项第一行。\n第二行。\n2. 第二项。\n3. 第三项。\n",
        encoding="utf-8",
    )
    assert module.parse_claims(compact_claims) == [
        "第一项第一行。 第二行。",
        "第二项。",
        "第三项。",
    ]

    abstract = tmp_path / "说明书摘要.md"
    abstract.write_text("# 示例装置\n\n这是摘要。\n\n摘要附图：图2。\n", encoding="utf-8")
    assert module.parse_abstract(abstract) == ("这是摘要。", 2)

    drawing_dir = tmp_path / "说明书附图"
    drawing_dir.mkdir()
    for number in (1, 2):
        (drawing_dir / f"图{number}.png").write_bytes(b"png-placeholder")
    index = tmp_path / "说明书附图.md"
    index.write_text(
        """# 说明书附图

## 图1 系统结构示意图

![图1](说明书附图/图1.png)

## 图2 方法流程图

![图2](说明书附图/图2.png)
""",
        encoding="utf-8",
    )
    figures = module.parse_figures(index)
    assert [(figure.number, figure.title) for figure in figures] == [
        (1, "系统结构示意图"),
        (2, "方法流程图"),
    ]


def test_direct_omml_builds_native_fraction_subscript_and_superscript(tmp_path):
    module = ASSEMBLER
    formula = "C(P,a_i)+(x^2)/(y_1-z)"
    root = module.build_omath(formula)
    xml = etree.tostring(root)
    namespaces = {"m": "http://schemas.openxmlformats.org/officeDocument/2006/math"}
    parsed = etree.fromstring(xml)
    assert len(parsed.xpath(".//m:f", namespaces=namespaces)) == 1
    assert len(parsed.xpath(".//m:sSub", namespaces=namespaces)) == 2
    assert len(parsed.xpath(".//m:sSup", namespaces=namespaces)) == 1

    document = Document()
    paragraph = document.add_paragraph("公式：")
    registry = module.MathRegistry()
    module.append_omath(paragraph, registry, formula)
    output = tmp_path / "native-math.docx"
    document.save(output)
    with ZipFile(output) as archive:
        document_xml = archive.read("word/document.xml")
    package_root = etree.fromstring(document_xml)
    assert len(package_root.xpath(".//m:oMath", namespaces=namespaces)) == 1
    assert b"[[EQ" not in document_xml


def test_linux_native_formula_no_longer_requires_word(monkeypatch, tmp_path):
    module = ASSEMBLER
    monkeypatch.setattr(module.sys, "platform", "linux")
    document = Document()
    paragraph = document.add_paragraph()
    registry = module.MathRegistry()
    module.append_omath(paragraph, registry, "ED_CT=Q_CT×C_CT(P,a)")
    output = tmp_path / "native-math.docx"
    document.save(output)
    assert module.process_with_word(output, registry, None) is None
    package = module.inspect_package(output, 1, "Heading1", "Strong")
    assert package["math_count"] == 1


@pytest.mark.skipif(
    shutil.which("libreoffice") is None or shutil.which("pdfinfo") is None,
    reason="需要 LibreOffice 和 pdfinfo 验证 Linux 公式渲染",
)
def test_linux_omml_roundtrip_renders_to_pdf(tmp_path):
    module = ASSEMBLER
    document = Document()
    paragraph = document.add_paragraph("公式：")
    registry = module.MathRegistry()
    module.append_omath(paragraph, registry, "C(P,a_i)+(x^2)/(y_1-z)")
    output = tmp_path / "native-math.docx"
    pdf = tmp_path / "native-math.pdf"
    document.save(output)
    module.export_pdf_with_libreoffice(output, pdf)
    assert pdf.is_file() and pdf.stat().st_size > 0
    assert module.count_pdf_pages(pdf) == 1


def test_linux_build_document_generates_native_math_without_word(monkeypatch, tmp_path):
    module = ASSEMBLER
    monkeypatch.setattr(module.sys, "platform", "linux")
    template = tmp_path / "输出模版.docx"
    _write_test_template(template)
    source = tmp_path / "02-申请文件"
    figures = source / "说明书附图"
    figures.mkdir(parents=True)
    (source / "权利要求书.md").write_text(
        "# 权利要求书\n\n1. 一种装置，包括参数 $a_i$。\n", encoding="utf-8"
    )
    (source / "说明书.md").write_text(
        "# 示例装置\n\n## 技术领域\n\n涉及参数 $a_i$。\n\n"
        "## 背景技术\n\n现有方案不足。\n\n## 发明内容\n\n"
        "$$\nC(P,a)=C(P,a_i)+(a-a_i)×[C(P,a_(i+1))-C(P,a_i)]/(a_(i+1)-a_i)\n$$\n\n"
        "## 附图说明\n\n图1是结构图。\n\n## 具体实施方式\n\n执行计算。\n",
        encoding="utf-8",
    )
    (source / "说明书摘要.md").write_text(
        "# 示例装置\n\n本发明提供一种示例装置。\n\n摘要附图：图1。\n",
        encoding="utf-8",
    )
    Image.new("RGB", (200, 100), "white").save(figures / "图1.png")
    (source / "说明书附图.md").write_text(
        "# 说明书附图\n\n## 图1 结构图\n\n![图1](说明书附图/图1.png)\n",
        encoding="utf-8",
    )
    output = tmp_path / "申请文件.docx"
    report = module.build_document(source, template, output, None)
    assert report["status"] == "STRUCTURE_VERIFIED"
    assert report["math_generation"] == {
        "engine": "direct_omml",
        "platform": "linux",
        "word_automation_required": False,
    }
    assert report["counts"]["native_word_math_objects"] == 3
    assert report["counts"]["pages"] is None
    with ZipFile(output) as archive:
        xml = archive.read("word/document.xml")
    assert xml.count(b"<m:oMath>") == 3
    assert b"<m:f>" in xml
    assert b"<m:sSub>" in xml


def test_visual_review_is_opt_in(monkeypatch, tmp_path):
    module = ASSEMBLER
    monkeypatch.setattr(
        module.sys,
        "argv",
        ["assemble_application_docx.py", "--case-dir", str(tmp_path)],
    )
    args = module.parse_args()
    assert args.visual_review is False

    monkeypatch.setattr(
        module.sys,
        "argv",
        [
            "assemble_application_docx.py",
            "--case-dir",
            str(tmp_path),
            "--visual-review",
        ],
    )
    args = module.parse_args()
    assert args.visual_review is True
